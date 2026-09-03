from copy import deepcopy
from pathlib import Path
import shutil
import stat

from docx import Document
from docx.oxml.ns import qn


REFERENCE = Path(r"D:\tools\xwechat_files\wxid_tv257izcdtfg11_875e\msg\file\2026-08\熊仿杰7.27(1).docx")
OUTPUT = Path(r"D:\code\GroupProject\Rafaels-Bikers-Game-Engine\output\徐华芳个案护理.docx")


def unique_cells(row):
    cells = []
    seen = set()
    for cell in row.cells:
        marker = id(cell._tc)
        if marker not in seen:
            seen.add(marker)
            cells.append(cell)
    return cells


def replace_paragraph_text(paragraph, text, bold_first_line=False):
    base_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        base_rpr = deepcopy(paragraph.runs[0]._r.rPr)

    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    lines = text.split("\n")
    for index, line in enumerate(lines):
        run = paragraph.add_run()
        if base_rpr is not None:
            run._r.insert(0, deepcopy(base_rpr))
        if index:
            run.add_break()
        run.add_text(line)
        if bold_first_line and index == 0:
            run.bold = True


def set_cell_text(cell, text, bold_first_line=False):
    paragraphs = cell.paragraphs
    first = paragraphs[0]
    for paragraph in paragraphs[1:]:
        cell._tc.remove(paragraph._p)
    replace_paragraph_text(first, text, bold_first_line=bold_first_line)


def set_right_cell(table, row_index, text, bold_first_line=False):
    set_cell_text(unique_cells(table.rows[row_index])[-1], text, bold_first_line)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
if OUTPUT.exists():
    OUTPUT.chmod(stat.S_IREAD | stat.S_IWRITE)
shutil.copy2(REFERENCE, OUTPUT)
OUTPUT.chmod(stat.S_IREAD | stat.S_IWRITE)
document = Document(OUTPUT)
table = document.tables[0]

# Title and patient identification.
replace_paragraph_text(document.paragraphs[0], "\t徐华芳的个案护理")

row0 = unique_cells(table.rows[0])
set_cell_text(row0[1], "65病区")
set_cell_text(row0[3], "102")
set_cell_text(row0[5], "徐华芳")

row1 = unique_cells(table.rows[1])
set_cell_text(row1[1], "女")
set_cell_text(row1[3], "93")
set_cell_text(row1[5], "612290")

set_right_cell(
    table,
    2,
    "高血压病3级（极高危）；帕金森病；认知障碍；阵发性心房颤动；室性期前收缩；房性期前收缩（房性早搏）；脑梗死后遗症；营养不良；低蛋白血症；胸腔积液；全身炎症反应综合征；急性肾功能不全",
)

set_right_cell(
    table,
    4,
    "患者徐华芳，女，93岁，因“高血压病50余年，控制欠佳1周”于2026/08/05入院。既往血压最高约200/100 mmHg，长期根据血压波动调整降压方案。\n住院后患者痰液增多，胸部影像提示肺部渗出较前增多；痰培养及药敏后先后予抗感染、抗真菌及化痰等治疗。2026/08/17出现少尿及代谢性酸中毒，随后肌酐升高。2026/08/22突发低血压（最低约83/45 mmHg），短期予去甲肾上腺素维持血压后转监护室进一步治疗。\n监护室期间予心电监护、气管切开接面罩吸氧及BiPAP辅助通气交替、抗感染/抗凝、营养支持及维持水电解质平衡等治疗，炎症指标好转后于2026/08/27转入65病区。现仍病危、一级护理，呼之可睁眼但无指令性动作，继续气道管理、营养支持及严密监测。",
)

set_right_cell(
    table,
    5,
    "疾病史：\n1. 高血压病史50余年，血压波动明显，长期口服降压药。\n2. 帕金森病、认知障碍、脑梗死后遗症。\n3. 阵发性心房颤动、室性期前收缩、房性期前收缩。\n4. 骨质疏松、骨关节病；右侧陈旧性股骨颈及股骨粗隆间骨折，2022年曾行手术治疗。\n5. 左侧肝占位性病变、肿瘤标记物升高、椎管内肿物。\n6. 胃术后及胆囊切除术后状态，长期气管造口。\n7. 焦虑状态；药物过敏史：青霉素。",
)

set_right_cell(
    table,
    6,
    "1. 2026/08/17血气分析：pH 7.28，PaCO₂ 28 mmHg，PaO₂ 226 mmHg，实际碳酸氢盐13.1 mmol/L，标准碳酸氢盐15 mmol/L，剩余碱-12.1 mmol/L，提示代谢性酸中毒。\n2. 住院期间床旁胸片提示肺部渗出较前增多；痰培养提示白假丝酵母菌阳性。\n3. 2026/08/21床旁超声提示右侧枕部皮下水肿。\n4. 监护室气道检查提示声门上组织水肿。\n5. 2026/08/28心肌标志物：BNP 6091 pg/mL，心肌肌钙蛋白T 0.071 ng/mL升高。",
)

set_right_cell(
    table,
    7,
    "2026/08/27-08/28主要异常指标：\n1. 血红蛋白96 g/L、红细胞计数3.40-3.62×10¹²/L、红细胞压积29.3%-31.5%，提示贫血；血小板119×10⁹/L降低。\n2. 淋巴细胞计数0.47-0.51×10⁹/L、淋巴细胞百分比11.2%-11.5%，均降低。\n3. 肌酐187 μmol/L、尿素13.7 mmol/L升高，eGFR 20 mL/min/1.73m²降低；二氧化碳20 mmol/L、血钙1.99 mmol/L降低。\n4. CRP 5.3 mg/L升高。\n5. D-二聚体由0.95升至2.11 mg/L，FDP 6.55 μg/mL升高；APTT、PT轻度延长。\n6. BNP 6091 pg/mL、心肌肌钙蛋白T 0.071 ng/mL升高。",
)

set_right_cell(
    table,
    9,
    "1. 病危、一级护理，持续心电监护，动态监测生命体征及血压变化。\n2. 气管切开接面罩吸氧与BiPAP辅助通气交替；按需湿化、雾化、吸痰，维持气道通畅。\n3. 记录24小时出入量，动态监测肾功能、电解质及酸碱平衡，遵医嘱利尿、补钙、补钠并调整液体治疗。\n4. 遵医嘱使用贝米肝素钠抗凝，结合血小板、凝血指标及痰中带血情况动态调整。\n5. 鼻肠管给予肠内营养（百普素0.5袋+温水500 mL，qd）并配合静脉营养、白蛋白等营养支持；每日鼻饲水约400 mL。\n6. 遵医嘱给予降压、抗心律失常、抗帕金森、化痰、抑酸、雾化吸入及对症支持治疗。",
)

set_right_cell(
    table,
    11,
    "2026/08/28 10:00：T 36.5℃，P 60，R 12，BP 137/55 mmHg，SpO₂ 100%。\n12:00：P 63，R 12，BP 133/63 mmHg，SpO₂ 100%。\n心律可不齐；静息NRS疼痛评分0分。",
)

set_right_cell(
    table,
    12,
    "伤口/皮肤：全身水肿；骶尾部约5×5 cm陈旧性红白色瘢痕，双侧髋部各约3×3 cm红白色瘢痕并有散在皮损，骨突处及骶尾部予泡沫敷料保护；后颈部肿胀。\n\n目前置管：\n1. 气管切开套管：接面罩吸氧，气囊充气，按需湿化吸痰；可吸出黄白色或白色黏痰，量较多。\n2. 左侧CVC：固定妥，无红肿、无外渗。\n3. 鼻肠管：固定妥，用于肠内营养及鼻饲给药。\n4. 留置导尿管：固定妥，尿液黄色，持续观察尿量及性状。",
)

set_right_cell(
    table,
    13,
    "卧位：卧床，取抬高床头及防误吸体位；四肢水肿，适当抬高。\n活动：目前无自主活动能力，协助每2小时翻身，进行被动关节活动及肢体功能维护。",
)

set_right_cell(table, 14, "ADL生活活动能力量表5分，为重度依赖；生活照护需由护理人员及照护者完成。")
set_right_cell(table, 15, "跌倒风险临床判定18分，为高风险；卧床期间落实床栏、警示标识及转运陪护。")
set_right_cell(table, 16, "Braden评分12分，为压力性损伤高度危险；Cubbin-Jackson评分30分，为低风险。结合现有水肿及皮损，按高风险落实气垫床、翻身减压和皮肤观察。")
set_right_cell(table, 17, "Padua静脉血栓风险评分9分，为高危；卧床、高龄及活动受限，需落实机械/药物预防并观察出血及下肢血栓征象。")
set_right_cell(table, 18, "NRS-2002评分6分，有营养风险；存在营养不良、低蛋白血症及吞咽障碍，给予鼻肠管肠内营养并结合静脉营养支持。")
set_right_cell(table, 19, "医院获得性肺炎风险评分19分，为超高危险；气管切开、痰液多且咳痰能力差，需加强气道管理及感染监测。")
set_right_cell(table, 20, "目前无新发手术并发症。重点观察气管造口、CVC、鼻肠管及导尿管相关并发症，并防范误吸、肺部感染、出血、血栓及压伤。")

set_right_cell(
    table,
    22,
    "1. 气体交换障碍 —— 与气管切开、痰液潴留及肺部感染有关\n护理措施：\n1. 持续观察呼吸频率、节律、SpO₂及缺氧表现，异常及时报告。\n2. 取抬高床头/半卧位，按医嘱面罩吸氧与BiPAP辅助通气交替。\n3. 保持气道湿化，按需吸痰、叩背，记录痰液颜色、量和性状。\n4. 做好气管切开护理，监测气囊压力，严格无菌操作并预防呼吸机相关感染。",
    bold_first_line=True,
)

set_right_cell(
    table,
    23,
    "2. 营养失调：低于机体需要量 —— 与吞咽障碍、长期卧床及低蛋白血症有关\n护理措施：\n1. 依医嘱经鼻肠管给予肠内营养，输注时抬高床头30°-45°，输注后保持体位30分钟以上。\n2. 观察腹胀、腹泻、反流及误吸征象，记录摄入量和排便情况。\n3. 定期监测体重、血红蛋白、白蛋白、电解质及肾功能，按需联系营养师。\n4. 做好口腔护理和鼻肠管维护，确保管路固定与通畅。",
    bold_first_line=True,
)

set_right_cell(
    table,
    24,
    "1. 误吸风险 —— 与意识障碍、吞咽障碍、气管切开及鼻肠管喂养有关\n护理措施：\n1. 持续采取防误吸体位，床头抬高30°-45°，翻身及操作时保护气道。\n2. 喂养前确认鼻肠管位置与通畅，严格控制速度和单次量。\n3. 观察呛咳、呼吸困难、SpO₂下降及痰液变化，必要时暂停喂养并通知医生。\n4. 床旁备吸引装置，按需行声门下分泌物吸引及气管切开吸痰。",
    bold_first_line=True,
)

set_right_cell(
    table,
    25,
    "2. 静脉血栓形成风险及皮肤完整性受损风险 —— 与高龄、长期卧床、水肿、营养不良和活动受限有关\n护理措施：\n1. 每2小时翻身，使用气垫床和减压敷料，保持皮肤清洁干燥，重点观察骶尾部、髋部及骨突处。\n2. 抬高水肿肢体，进行被动踝泵及关节活动，观察下肢周径、皮温、颜色、疼痛和肿胀。\n3. 遵医嘱落实抗凝治疗，监测血小板、凝血指标、皮肤黏膜及痰液有无出血。\n4. 严格维护CVC和导尿管，落实导管相关感染预防。",
    bold_first_line=True,
)

set_right_cell(
    table,
    27,
    "1. 目前以鼻肠管肠内营养为主，严格执行医嘱配方、速度及冲管水量，不自行经口喂食。\n2. 营养输注时及输注后保持床头抬高，观察反流、呛咳、腹胀、腹泻等情况。\n3. 每日记录摄入量、排便及体重变化；出现喂养不耐受及时告知医护人员。",
)

set_right_cell(
    table,
    28,
    "皮肤方面：\n1. 每2小时协助翻身，避免拖、拉、推；保持骶尾部、髋部及会阴部清洁干燥。\n2. 观察瘢痕、破损及水肿部位，敷料污染、松脱或渗液时及时更换并报告。\n导管方面：\n1. 向家属说明气管切开、CVC、鼻肠管及导尿管的用途，勿自行牵拉或拔除。\n2. 保持各管路固定、通畅，出现红肿、渗液、堵塞、脱出或尿量异常时及时通知医护人员。",
)

set_right_cell(
    table,
    29,
    "1. 青霉素过敏，任何诊疗环节均应主动告知医护人员。\n2. 抗凝药：观察皮肤黏膜、痰液、尿液及大便有无出血，出现异常立即报告。\n3. 降压及抗心律失常药：按时服用，密切监测血压、心率，不自行停药或调量。\n4. 抗帕金森及其他药物须按医嘱经鼻肠管给药，注意嗜睡、低血压等不良反应。",
)

set_right_cell(
    table,
    30,
    "卧床协助翻身及肢体被动运动；转运时保护气切和导管，异常立即停止并报告。",
)

# Remove the empty intermediate section-break paragraph that generated a blank final page.
body = document._element.body
for paragraph_element in list(body.findall(qn("w:p"))):
    sect_pr = paragraph_element.find("./w:pPr/w:sectPr", namespaces=paragraph_element.nsmap)
    text_nodes = paragraph_element.findall(".//w:t", namespaces=paragraph_element.nsmap)
    if sect_pr is not None and not "".join(node.text or "" for node in text_nodes).strip():
        body.remove(paragraph_element)

document.save(OUTPUT)
print(OUTPUT)
